import pytest

from mail_merge.excel import read_preview, read_recipients


class TestReadRecipients:
    def test_basic(self, sample_xlsx):
        recipients = read_recipients(sample_xlsx, "email")
        assert len(recipients) == 2
        assert recipients[0]["name"] == "Alice"
        assert recipients[0]["email"] == "alice@example.com"
        assert recipients[1]["name"] == "Bob"

    def test_skips_empty_email(self, sample_xlsx):
        recipients = read_recipients(sample_xlsx, "email")
        emails = [r["email"] for r in recipients]
        assert "" not in emails
        assert len(recipients) == 2  # Charlie skipped

    def test_missing_email_column(self, sample_xlsx):
        with pytest.raises(ValueError, match="not found"):
            read_recipients(sample_xlsx, "nonexistent")

    def test_case_insensitive_column(self, sample_xlsx):
        recipients = read_recipients(sample_xlsx, "Email")
        assert len(recipients) == 2

    def test_empty_spreadsheet(self, tmp_path):
        import openpyxl
        path = tmp_path / "empty.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        # Remove default data — leave truly empty
        wb.save(path)

        with pytest.raises(ValueError, match="empty"):
            read_recipients(path, "email")

    def test_ragged_rows_padded(self, tmp_path):
        """Rows shorter than headers are padded with empty strings."""
        import openpyxl
        path = tmp_path / "ragged.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "email", "company"])
        ws.append(["Alice", "alice@example.com"])  # missing company
        wb.save(path)

        recipients = read_recipients(path, "email")
        assert len(recipients) == 1
        assert recipients[0]["name"] == "Alice"
        assert recipients[0]["company"] == ""

    def test_oversized_file_rejected(self, tmp_path, monkeypatch):
        """Files larger than 50 MB are rejected before loading."""
        import openpyxl
        path = tmp_path / "big.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["email"])
        wb.save(path)

        # Fake a large file size without writing 50 MB of data
        monkeypatch.setattr("pathlib.Path.stat", lambda self: type("S", (), {"st_size": 51 * 1024 * 1024})())

        with pytest.raises(ValueError, match="too large"):
            read_recipients(path, "email")


class TestReadPreview:
    def test_basic(self, sample_xlsx):
        columns, rows, sheets, total_rows, active_sheet = read_preview(sample_xlsx)
        assert "name" in columns
        assert "email" in columns
        assert total_rows == 3  # Alice, Bob, Charlie
        assert len(rows) == 3
        assert active_sheet == "Sheet"

    def test_truncation(self, tmp_path):
        """read_preview with max_rows limits returned rows."""
        import openpyxl
        path = tmp_path / "many.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "email"])
        for i in range(20):
            ws.append([f"User{i}", f"user{i}@example.com"])
        wb.save(path)

        columns, rows, sheets, total_rows, _active = read_preview(path, max_rows=5)
        assert total_rows == 20
        assert len(rows) == 5
        assert rows[0]["name"] == "User0"

    def test_ragged_rows_padded(self, tmp_path):
        """Preview rows shorter than headers are padded with empty strings."""
        import openpyxl
        path = tmp_path / "ragged.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "email", "company"])
        ws.append(["Alice", "alice@example.com"])  # missing company
        wb.save(path)

        columns, rows, sheets, total_rows, _active = read_preview(path)
        assert len(rows) == 1
        assert rows[0]["company"] == ""

    def test_empty_spreadsheet_raises(self, tmp_path):
        import openpyxl
        path = tmp_path / "empty.xlsx"
        wb = openpyxl.Workbook()
        wb.save(path)

        with pytest.raises(ValueError, match="empty"):
            read_preview(path)

    def test_oversized_file_rejected(self, tmp_path, monkeypatch):
        import openpyxl
        path = tmp_path / "big.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["email"])
        wb.save(path)

        monkeypatch.setattr(
            "pathlib.Path.stat",
            lambda self: type("S", (), {"st_size": 51 * 1024 * 1024})(),
        )
        with pytest.raises(ValueError, match="too large"):
            read_preview(path)
