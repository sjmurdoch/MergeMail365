import logging
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

_MAX_SPREADSHEET_SIZE = 50 * 1024 * 1024  # 50 MB


def _check_file_size(path: Path) -> None:
    try:
        file_size = path.stat().st_size
    except FileNotFoundError:
        raise FileNotFoundError(f"Spreadsheet not found: {path}")
    if file_size > _MAX_SPREADSHEET_SIZE:
        raise ValueError(
            f"Spreadsheet too large: {file_size / (1024 * 1024):.0f} MB "
            f"(limit: {_MAX_SPREADSHEET_SIZE // (1024 * 1024)} MB)"
        )


def _cell_str(cell: object) -> str:
    """Convert a cell value to a stripped string, treating None as empty."""
    return str(cell.value).strip() if cell.value is not None else ""  # type: ignore[attr-defined]


def _parse_headers(header_row: tuple) -> list[str]:  # type: ignore[type-arg]
    return [_cell_str(cell) for cell in header_row]


def read_preview(
    path: str | Path,
    sheet_name: str | None = None,
    max_rows: int = 99,
) -> tuple[list[str], list[dict[str, str]], list[str], int, str]:
    """Read column headers, first N rows, sheet names, and total row count.

    Returns (columns, rows, sheet_names, total_rows, active_sheet).
    Does not require an email_column — used for spreadsheet preview in the web UI.
    """
    path = Path(path)
    _check_file_size(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        if sheet_name:
            ws = wb[sheet_name]
        else:
            ws = wb.active
            sheet_name = ws.title

        # Count total rows (excluding header)
        total_rows = ws.max_row - 1 if ws.max_row else 0

        rows_iter = ws.iter_rows()
        header_row = next(rows_iter, None)
        if header_row is None:
            raise ValueError("Spreadsheet is empty")

        headers = _parse_headers(header_row)
        columns = [h for h in headers if h]

        preview_rows: list[dict[str, str]] = []
        for _, row in zip(range(max_rows), rows_iter):
            values = [_cell_str(cell) for cell in row]
            while len(values) < len(headers):
                values.append("")
            record = {headers[i]: values[i] for i in range(len(headers)) if headers[i]}
            preview_rows.append(record)
    finally:
        wb.close()
    return columns, preview_rows, sheet_names, total_rows, sheet_name


def read_recipients(
    path: str | Path,
    email_column: str,
    sheet_name: str | None = None,
) -> list[dict[str, str]]:
    """Read recipients from an Excel spreadsheet.

    Returns a list of dicts mapping column header -> cell value.
    Rows with an empty email cell are skipped with a warning.
    Raises ValueError if the email column is not found.
    """
    path = Path(path)
    _check_file_size(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active

        rows = ws.iter_rows()
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Spreadsheet is empty")

        headers = _parse_headers(header_row)

        # Find the email column (case-insensitive)
        email_col_idx = None
        for i, h in enumerate(headers):
            if h.lower() == email_column.lower():
                email_col_idx = i
                break

        if email_col_idx is None:
            raise ValueError(
                f"Email column '{email_column}' not found. "
                f"Available columns: {', '.join(h for h in headers if h)}"
            )

        recipients = []
        for row_num, row in enumerate(rows, start=2):
            values = [_cell_str(cell) for cell in row]
            while len(values) < len(headers):
                values.append("")

            email_value = values[email_col_idx]
            if not email_value:
                if any(v for v in values):
                    logger.warning("Row %d: empty email, skipping", row_num)
                continue

            record = {headers[i]: values[i] for i in range(len(headers)) if headers[i]}
            recipients.append(record)
    finally:
        wb.close()
    return recipients
